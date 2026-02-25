"""
conftest.py — Fixtures compartilhadas entre toda a suite de testes.
"""
import sys
from pathlib import Path

# Garante que src/ está no path para imports relativos funcionarem
ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "src"
for p in (str(ROOT), str(SRC)):
    if p not in sys.path:
        sys.path.insert(0, p)

import io
import datetime
import pytest
import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------------------
# Helpers para construir xlsx em memória
# ---------------------------------------------------------------------------

def _make_xlsx_bytes(sheet_data: dict[str, list[list]]) -> bytes:
    """
    Cria um arquivo .xlsx em memória com as abas fornecidas.
    sheet_data: { "NomeAba": [[linha0col0, ...], [linha1col0, ...], ...] }
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove aba padrão
    for sheet_name, rows in sheet_data.items():
        ws = wb.create_sheet(title=sheet_name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def _write_tmp_xlsx(tmp_path: Path, sheet_data: dict[str, list[list]], filename="test.xlsx") -> Path:
    """Grava o xlsx em disco e retorna o caminho."""
    path = tmp_path / filename
    path.write_bytes(_make_xlsx_bytes(sheet_data))
    return path


# ---------------------------------------------------------------------------
# Fixtures de planilha Layout padrão
# ---------------------------------------------------------------------------

BASE_DATE = datetime.date(2025, 7, 1)


def _layout_rows(doc="AZ", plano="1.01.02.01", n_rows=3, valor_base=1000.0, offset_days=0):
    """Gera linhas de um bloco Layout com metadados + cabeçalho + dados."""
    rows = [
        [doc, plano],          # metadados (linha 0)
        ["Contrato", "Valor", "Data Crédito"],  # cabeçalho (linha 1)
    ]
    for i in range(n_rows):
        d = BASE_DATE + datetime.timedelta(days=i + offset_days)
        rows.append([f"CTR-{i+1:03d}", valor_base + i * 100, d])
    return rows


@pytest.fixture()
def layout_simples_xlsx(tmp_path):
    """Planilha com 1 bloco, 3 linhas de dados."""
    rows = _layout_rows()
    return _write_tmp_xlsx(tmp_path, {"Layout": rows})


@pytest.fixture()
def layout_multiplos_blocos_xlsx(tmp_path):
    """Planilha com 2 blocos horizontais (6 colunas).

    Estrutura esperada pelo leitor:
      linha 0: [AZ, 1.01.02.01, , REG, 1.04.01.08, ]   <- metadados de ambos
      linha 1: [Contrato, Valor, Data Crédito, Contrato, Valor, Data Crédito]  <- cabeçalho
      linha 2+: dados
    """
    n1, n2 = 3, 2
    dados1 = [
        [f"CTR-{i+1:03d}", 1000.0 + i * 100, BASE_DATE + datetime.timedelta(days=i)]
        for i in range(n1)
    ]
    dados2 = [
        [f"REG-{i+1:03d}", 500.0 + i * 100, BASE_DATE + datetime.timedelta(days=i + 5)]
        for i in range(n2)
    ]

    max_data = max(n1, n2)
    combined = []
    # linha de metadados
    combined.append(["AZ", "1.01.02.01", None, "REG", "1.04.01.08", None])
    # linha de cabeçalho
    combined.append(["Contrato", "Valor", "Data Crédito", "Contrato", "Valor", "Data Crédito"])
    # linhas de dados
    for i in range(max_data):
        r1 = dados1[i] if i < n1 else [None, None, None]
        r2 = dados2[i] if i < n2 else [None, None, None]
        combined.append(r1 + r2)

    return _write_tmp_xlsx(tmp_path, {"Layout": combined})


@pytest.fixture()
def layout_valor_zerado_xlsx(tmp_path):
    """Planilha com 1 bloco contendo linhas de valor zero (devem ser ignoradas)."""
    rows = [
        ["AZ", "1.01.02.01"],
        ["Contrato", "Valor", "Data Crédito"],
        ["CTR-001", 1000.0, BASE_DATE],
        ["CTR-002", 0.0, BASE_DATE + datetime.timedelta(days=1)],
        ["CTR-003", 2000.0, BASE_DATE + datetime.timedelta(days=2)],
    ]
    return _write_tmp_xlsx(tmp_path, {"Layout": rows})


@pytest.fixture()
def layout_sem_cabecalho_xlsx(tmp_path):
    """Planilha sem linha de cabeçalho válida."""
    rows = [
        ["AZ", "1.01.02.01"],
        ["ColA", "ColB", "ColC"],
        ["CTR-001", 1000.0, BASE_DATE],
    ]
    return _write_tmp_xlsx(tmp_path, {"Layout": rows})


@pytest.fixture()
def layout_valor_formato_br_xlsx(tmp_path):
    """Planilha com valores em formato BR (1.234,56) como strings."""
    rows = [
        ["AZ", "1.01.02.01"],
        ["Contrato", "Valor", "Data Crédito"],
        ["CTR-001", "1.234,56", BASE_DATE],
        ["CTR-002", "628,91", BASE_DATE + datetime.timedelta(days=1)],
        ["CTR-003", "(500,00)", BASE_DATE + datetime.timedelta(days=2)],  # negativo
    ]
    return _write_tmp_xlsx(tmp_path, {"Layout": rows})


# ---------------------------------------------------------------------------
# DataFrames prontos para testes de processador / gerador
# ---------------------------------------------------------------------------

@pytest.fixture()
def df_bloco_simples():
    """DataFrame representando um bloco já lido."""
    return pd.DataFrame({
        "Contrato": ["CTR-001", "CTR-002", "CTR-003"],
        "Valor": [1000.0, 1100.0, 1200.0],
        "Data Crédito": pd.to_datetime([
            "2025-07-01", "2025-07-02", "2025-07-03"
        ]),
    })


@pytest.fixture()
def dados_por_bloco_fixture(df_bloco_simples):
    """Dict {(doc, plano): DataFrame} pronto para processador."""
    df2 = pd.DataFrame({
        "Contrato": ["CTR-A", "CTR-B"],
        "Valor": [500.0, 600.0],
        "Data Crédito": pd.to_datetime(["2025-07-05", "2025-07-06"]),
    })
    return {
        ("AZ", "1.01.02.01"): df_bloco_simples.copy(),
        ("REG", "1.04.01.08"): df2,
    }
